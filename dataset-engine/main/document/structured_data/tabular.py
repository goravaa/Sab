
"""Read tabular files from S3-compatible storage and yield Documents.

Supported formats:
    - CSV
    - XLSX
    - XLSM

Responsibilities:
    1. Read source objects through S3Connector.
    2. Stream rows without materializing the entire dataset in RAM.
    3. Convert rows into the existing Document model.
    4. Generate deterministic document IDs and content signatures.

Chunking, embedding, and storage belong to downstream pipeline stages.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import tempfile

from collections.abc import Iterator, Mapping, Sequence
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any
from uuid import UUID


from main.schema.document import Document, SourceType
from main.connect.object_storage.s3 import S3Connector, SourceObject


SUPPORTED_EXTENSIONS = frozenset({".csv", ".xlsx", ".xlsm"})
SIGNATURE_VERSION = 1


# ---------------------------------------------------------------------------
# Serialization
# ---------------------------------------------------------------------------

def canonicalize(value: Any) -> Any:
    """Convert common tabular values into deterministic JSON-compatible data.

    This representation is used for document content, metadata, and hashing.
    """

    if value is None:
        return None

    if isinstance(value, (str, bool, int)):
        return value

    if isinstance(value, float):
        if math.isnan(value):
            return None
        if not math.isfinite(value):
            return {"$float": str(value)}
        return value

    if isinstance(value, Decimal):
        if not value.is_finite():
            return {"$decimal": str(value)}
        return {"$decimal": str(value)}

    if isinstance(value, datetime):
        return {"$datetime": value.isoformat()}

    if isinstance(value, date):
        return {"$date": value.isoformat()}

    if isinstance(value, time):
        return {"$time": value.isoformat()}

    if isinstance(value, UUID):
        return {"$uuid": str(value)}

    if isinstance(value, bytes):
        return {"$bytes": value.hex()}

    if isinstance(value, Mapping):
        return {
            str(key): canonicalize(item)
            for key, item in sorted(
                value.items(),
                key=lambda pair: str(pair[0]),
            )
        }

    if isinstance(value, (list, tuple)):
        return [canonicalize(item) for item in value]

    # Handle scalar missing-value types such as pandas.NA without
    # making pandas a required dependency.
    if type(value).__name__ in {"NAType", "NaTType"}:
        return None

    # Fail rather than silently hash an unknown object using an
    # unstable or lossy string representation.
    raise TypeError(
        f"Unsupported tabular value type: {type(value).__qualname__}"
    )


def canonical_json(value: Any) -> str:
    """Serialize a value deterministically."""

    return json.dumps(
        canonicalize(value),
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=False,
        allow_nan=False,
    )


def sha256_json(value: Any) -> str:
    """Return a SHA-256 hex digest of a canonical JSON representation."""

    return hashlib.sha256(
        canonical_json(value).encode("utf-8")
    ).hexdigest()


# ---------------------------------------------------------------------------
# Header handling
# ---------------------------------------------------------------------------

def normalize_headers(headers: Sequence[Any]) -> list[str]:
    """Normalize missing and duplicate column names deterministically."""

    normalized: list[str] = []
    counts: dict[str, int] = {}

    for index, value in enumerate(headers, start=1):
        name = str(value).strip() if value is not None else ""

        if not name:
            name = f"column_{index}"

        counts[name] = counts.get(name, 0) + 1

        if counts[name] > 1:
            name = f"{name}_{counts[name]}"

        normalized.append(name)

    return normalized


def make_row(
    headers: Sequence[str],
    values: Sequence[Any],
) -> dict[str, Any]:
    """Build a row mapping, filling missing trailing cells with None."""

    return {
        header: values[index] if index < len(values) else None
        for index, header in enumerate(headers)
    }


# ---------------------------------------------------------------------------
# Document construction
# ---------------------------------------------------------------------------

def row_to_document(
    row: Mapping[str, Any],
    *,
    namespace: str,
    source_id: str,
    row_number: int,
    sheet_name: str | None = None,
    id_column: str | None = None,
) -> Document:
    """Convert one tabular row into the project's Document model.

    If id_column is configured and populated, its value identifies the row.
    Otherwise, the physical row number is used.

    Row-number identities are not stable across insertions or reordering.
    """

    values = {
        str(column): canonicalize(value)
        for column, value in row.items()
    }

    if id_column is not None and id_column not in values:
        raise ValueError(
            f"Configured ID column {id_column!r} was not found. "
            f"Available columns: {list(values)}"
        )

    id_value = values.get(id_column) if id_column is not None else None

    if id_value is not None and id_value != "":
        row_identity = {
            "type": "column",
            "column": id_column,
            "value": id_value,
        }
    else:
        row_identity = {
            "type": "row_number",
            "value": row_number,
        }

    identity = {
        "namespace": namespace,
        "source_id": source_id,
        "sheet": sheet_name,
        "row_identity": row_identity,
    }

    # Stable for the same logical source and row identity.
    doc_id = sha256_json({
        "version": 1,
        **identity,
    })

    # Keep content rendering deterministic and readable.
    raw_content = "\n".join(
        f"{column}: {canonical_json(value)}"
        for column, value in values.items()
    )

    # The signature tracks the actual row content, not its identity.
    signature = sha256_json({
        "version": SIGNATURE_VERSION,
        "columns": list(values.keys()),
        "values": values,
    })

    return Document(
        doc_id=doc_id,
        source_type=SourceType.S3_OBJECT,
        raw_content=raw_content,
        signature=signature,
        metadata={
            **identity,
            "row_number": row_number,
            "row": values,
        },
    )


# ---------------------------------------------------------------------------
# CSV reader
# ---------------------------------------------------------------------------

def iter_csv_documents(
    connector: S3Connector,
    source: SourceObject,
    *,
    namespace: str = "files",
    id_column: str | None = None,
    encoding: str = "utf-8-sig",
    delimiter: str = ",",
) -> Iterator[Document]:
    """Stream CSV rows from S3 and yield Documents.

    Memory usage is bounded by the streaming reader and the current row.
    The first CSV record is treated as the header.
    """

    with connector.open_object(source.key) as body:
        text = io.TextIOWrapper(
            body,
            encoding=encoding,
            newline="",
        )

        try:
            reader = csv.reader(text, delimiter=delimiter)
            raw_headers = next(reader, None)

            if raw_headers is None:
                return

            headers = normalize_headers(raw_headers)

            for row_number, cells in enumerate(reader, start=2):
                if not cells or all(cell == "" for cell in cells):
                    continue

                row = make_row(headers, cells)

                yield row_to_document(
                    row,
                    namespace=namespace,
                    source_id=source.key,
                    row_number=row_number,
                    id_column=id_column,
                )
        finally:
            # The connector owns the underlying response body.
            # Detaching prevents TextIOWrapper from closing it early.
            text.detach()


# ---------------------------------------------------------------------------
# Excel reader
# ---------------------------------------------------------------------------

def iter_excel_documents(
    connector: S3Connector,
    source: SourceObject,
    *,
    namespace: str = "files",
    id_column: str | None = None,
) -> Iterator[Document]:
    """Read XLSX/XLSM from S3 and yield Documents sheet by sheet.

    XLSX is a ZIP-based format requiring seekable input. The object is
    streamed to a temporary disk file, then parsed in read-only mode.
    The complete workbook is not loaded into RAM.
    """

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Excel support requires openpyxl; install 'openpyxl'."
        ) from exc

    with connector.open_object(source.key) as body:
        with tempfile.TemporaryFile(suffix=".xlsx") as temporary_file:
            while True:
                chunk = body.read(8 * 1024 * 1024)

                if not chunk:
                    break

                temporary_file.write(chunk)

            temporary_file.seek(0)

            workbook = load_workbook(
                temporary_file,
                read_only=True,
                data_only=True,
            )

            try:
                for worksheet in workbook.worksheets:
                    rows = worksheet.iter_rows(values_only=True)
                    raw_headers = next(rows, None)

                    if raw_headers is None:
                        continue

                    headers = normalize_headers(raw_headers)

                    for row_number, cells in enumerate(rows, start=2):
                        if all(value is None for value in cells):
                            continue

                        row = make_row(headers, cells)

                        yield row_to_document(
                            row,
                            namespace=namespace,
                            source_id=source.key,
                            row_number=row_number,
                            sheet_name=worksheet.title,
                            id_column=id_column,
                        )
            finally:
                workbook.close()


# ---------------------------------------------------------------------------
# Source dispatch
# ---------------------------------------------------------------------------

def iter_tabular_documents(
    connector: S3Connector,
    source: SourceObject,
    *,
    namespace: str = "files",
    id_column: str | None = None,
    encoding: str = "utf-8-sig",
    delimiter: str = ",",
) -> Iterator[Document]:
    """Dispatch a source object to the appropriate tabular reader.

    Unsupported file formats raise ValueError instead of being silently
    ignored. Callers can filter source objects before invoking this function.
    """

    key = source.key.lower()
    extension = next(
        (
            ext
            for ext in SUPPORTED_EXTENSIONS
            if key.endswith(ext)
        ),
        None,
    )

    if extension is None:
        raise ValueError(
            f"Unsupported tabular file: {source.key!r}. "
            f"Supported extensions: {sorted(SUPPORTED_EXTENSIONS)}"
        )

    if extension == ".csv":
        yield from iter_csv_documents(
            connector,
            source,
            namespace=namespace,
            id_column=id_column,
            encoding=encoding,
            delimiter=delimiter,
        )
        return

    yield from iter_excel_documents(
        connector,
        source,
        namespace=namespace,
        id_column=id_column,
    )
