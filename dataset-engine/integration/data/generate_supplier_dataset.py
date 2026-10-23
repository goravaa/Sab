"""Generate a connected, synthetic B2B supplier dataset in Postgres and S3."""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import random
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

INTEGRATION_ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT))

import boto3
import psycopg
from botocore.exceptions import ClientError
from docx import Document as WordDocument
from openpyxl import Workbook
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from main.connect.database.postgres import PostgresConnector
from main.connect.object_storage.s3 import S3Connector
from main.schema.service import DatabaseConfig, S3Config


ROOT = INTEGRATION_ROOT
ENV_FILE = ROOT / "docker" / ".env"
OUTPUT_ROOT = ROOT / "generated"
DB_SCHEMA = "supplier_demo"
BUCKET = "sab-integration"

FIRST_NAMES = [
    "Avery", "Jordan", "Morgan", "Riley", "Casey", "Taylor", "Cameron", "Quinn",
    "Reese", "Parker", "Rowan", "Skyler", "Drew", "Emerson", "Finley", "Hayden",
    "Jamie", "Kendall", "Logan", "Micah", "Payton", "Sage", "Shiloh", "Tatum",
    "Alex", "Bailey", "Dakota", "Elliot", "Frankie", "Harper", "Jules", "Kai",
]
LAST_NAMES = [
    "Bennett", "Carter", "Diaz", "Ellis", "Foster", "Garcia", "Hayes", "Ibrahim",
    "Jensen", "Kim", "Lopez", "Miller", "Nguyen", "Owens", "Patel", "Quinn",
    "Reed", "Shah", "Turner", "Usman", "Vega", "Walker", "Xu", "Young",
    "Zimmerman", "Brooks", "Cole", "Dawson", "Evans", "Flores", "Gray", "Hughes",
]
COMPANY_A = [
    "Cedar", "Summit", "Harbor", "Northstar", "Bluebird", "Redwood", "Silverline",
    "Pioneer", "Evergreen", "Westbridge", "Clearwater", "Ironwood", "Brighton", "Keystone",
]
COMPANY_B = [
    "Supply", "Industrial", "Field Services", "Equipment", "Manufacturing", "Logistics",
    "Systems", "Outfitters", "Works", "Distribution", "Technologies", "Contracting",
]
CITIES = [
    ("Seattle", "WA", "98101"), ("Portland", "OR", "97201"), ("San Francisco", "CA", "94105"),
    ("Los Angeles", "CA", "90012"), ("San Diego", "CA", "92101"), ("Phoenix", "AZ", "85004"),
    ("Denver", "CO", "80202"), ("Austin", "TX", "78701"), ("Dallas", "TX", "75201"),
    ("Chicago", "IL", "60601"), ("Minneapolis", "MN", "55401"), ("Nashville", "TN", "37201"),
    ("Atlanta", "GA", "30303"), ("Miami", "FL", "33130"), ("Charlotte", "NC", "28202"),
    ("Columbus", "OH", "43215"), ("Detroit", "MI", "48226"), ("Boston", "MA", "02108"),
]
PRODUCTS = [
    ("Cordless Drill", "power tools", "189.00"), ("Impact Driver", "power tools", "219.00"),
    ("Circular Saw", "power tools", "249.00"), ("Angle Grinder", "power tools", "129.00"),
    ("Rotary Hammer", "power tools", "329.00"), ("Laser Measure", "measuring", "99.00"),
    ("Digital Caliper", "measuring", "54.00"), ("Torque Wrench", "hand tools", "78.00"),
    ("Socket Set", "hand tools", "64.00"), ("Work Light", "site equipment", "42.00"),
    ("Safety Helmet", "safety", "28.00"), ("Cut Resistant Gloves", "safety", "19.00"),
    ("Hearing Protection", "safety", "32.00"), ("Respirator Kit", "safety", "46.00"),
    ("Extension Cable 30m", "electrical", "58.00"), ("Portable Generator", "site equipment", "799.00"),
    ("Air Compressor", "site equipment", "549.00"), ("Water Pump", "site equipment", "285.00"),
    ("Steel Storage Cabinet", "storage", "389.00"), ("Parts Organizer", "storage", "36.00"),
    ("Hydraulic Jack", "lifting", "179.00"), ("Chain Hoist", "lifting", "265.00"),
    ("Pallet Truck", "material handling", "449.00"), ("Platform Cart", "material handling", "149.00"),
    ("Welding Helmet", "welding", "119.00"), ("MIG Welder", "welding", "689.00"),
    ("Shop Vacuum", "site equipment", "159.00"), ("Pressure Washer", "site equipment", "319.00"),
    ("Concrete Mixer", "site equipment", "429.00"), ("Survey Tripod", "measuring", "89.00"),
]
CARRIERS = ["ParcelWay", "Northstar Freight", "Pioneer Express", "MetroHaul"]
TICKET_CATEGORIES = ["delivery", "product question", "warranty", "replacement part", "billing"]
SUBJECTS = [
    "Delivery timing for a project order", "Request for a replacement part", "Warranty coverage question",
    "Invoice line item clarification", "Product specification confirmation", "Update delivery instructions",
    "Packaging arrived damaged", "Request a copy of the product manual",
]


def load_env(path: Path) -> None:
    if not path.exists():
        return
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if line and not line.startswith("#") and "=" in line:
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip("\"'"))


def need_env(name: str) -> str:
    value = os.getenv(name)
    if not value:
        raise RuntimeError(f"Missing {name}; configure it in {ENV_FILE}")
    return value


def money(value: Decimal | str | int | float) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def iso(value: date | datetime | None) -> str:
    return value.isoformat() if value is not None else ""


def company_name(rng: random.Random) -> str:
    return f"{rng.choice(COMPANY_A)} {rng.choice(COMPANY_B)} {rng.choice(['LLC', 'Inc.', 'Group', 'Co.'])}"


def build_data(seed: int, count: int) -> dict[str, Any]:
    rng = random.Random(seed)
    today = date(2026, 10, 22)
    earliest = date(2025, 1, 1)
    date_span = (today - earliest).days
    prefix = f"supplier-demo/seed-{seed:04d}/"

    customers: list[dict[str, Any]] = []
    for n in range(1, count + 1):
        customer_id = f"CUST-{seed:04d}-{n:06d}"
        first, last = rng.choice(FIRST_NAMES), rng.choice(LAST_NAMES)
        city, state, postal = rng.choice(CITIES)
        customers.append({
            "customer_id": customer_id,
            "company_name": company_name(rng),
            "contact_name": f"{first} {last}",
            "email": f"{first.lower()}.{last.lower()}.{n}@example.test",
            "phone": f"+1-555-{rng.randrange(200, 999):03d}-{rng.randrange(1000):04d}",
            "segment": rng.choices(["small business", "mid-market", "enterprise"], [55, 35, 10])[0],
            "city": city,
            "state": state,
            "postal_code": postal,
            "created_at": earliest + timedelta(days=rng.randrange(date_span + 1)),
            "agreement_key": f"{prefix}customers/{customer_id}/agreement.docx",
        })

    products: list[dict[str, Any]] = []
    for index in range(1, 81):
        name, category, base_price = rng.choice(PRODUCTS)
        sku = f"TOOL-{seed:04d}-{index:04d}"
        price = money(Decimal(base_price) * Decimal(str(rng.uniform(0.82, 1.24))))
        products.append({
            "sku": sku,
            "product_name": name,
            "category": category,
            "unit_price": price,
            "active": rng.random() > 0.04,
            "reorder_level": rng.randrange(8, 36),
            "on_hand": rng.randrange(0, 180),
        })

    orders: list[dict[str, Any]] = []
    order_items: list[dict[str, Any]] = []
    invoices: list[dict[str, Any]] = []
    shipments: list[dict[str, Any]] = []
    tickets: list[dict[str, Any]] = []
    order_counter = item_counter = shipment_counter = ticket_counter = 0
    for customer in customers:
        customer_order_ids: list[str] = []
        for _ in range(rng.choices([1, 2, 3, 4, 5], [18, 24, 30, 20, 8])[0]):
            order_counter += 1
            order_id = f"ORD-{seed:04d}-{order_counter:07d}"
            placed = earliest + timedelta(days=rng.randrange(date_span + 1))
            status = rng.choices(
                ["delivered", "shipped", "processing", "cancelled", "returned"],
                [54, 18, 14, 8, 6],
            )[0]
            chosen_products = rng.sample(products, rng.randint(1, 4))
            line_total = Decimal("0.00")
            line_items = []
            for product in chosen_products:
                item_counter += 1
                quantity = rng.randint(1, 8)
                unit_price = product["unit_price"]
                line_total += unit_price * quantity
                line_items.append({
                    "item_id": f"ITEM-{seed:04d}-{item_counter:08d}",
                    "order_id": order_id,
                    "sku": product["sku"],
                    "quantity": quantity,
                    "unit_price": unit_price,
                    "line_total": money(unit_price * quantity),
                })
            tax = money(line_total * Decimal("0.0825"))
            freight = Decimal("0.00") if line_total >= 1000 else Decimal("24.95")
            total = money(line_total + tax + freight)
            order = {
                "order_id": order_id,
                "customer_id": customer["customer_id"],
                "placed_on": placed,
                "status": status,
                "subtotal": money(line_total),
                "tax": tax,
                "shipping": freight,
                "total": total,
                "items": line_items,
            }
            orders.append(order)
            customer_order_ids.append(order_id)
            order_items.extend(line_items)

            invoice_id = f"INV-{seed:04d}-{order_counter:07d}"
            invoice_status = rng.choices(["paid", "open", "overdue", "void"], [68, 18, 9, 5])[0]
            invoice_key = f"{prefix}invoices/{invoice_id}.pdf"
            invoices.append({
                "invoice_id": invoice_id,
                "order_id": order_id,
                "issued_on": placed + timedelta(days=rng.randint(0, 2)),
                "due_on": placed + timedelta(days=30),
                "status": invoice_status,
                "total": total,
                "object_key": invoice_key,
                "items": line_items,
            })

            if status in {"delivered", "shipped", "returned"}:
                shipment_counter += 1
                shipped_on = placed + timedelta(days=rng.randint(1, 5))
                delivered_on = shipped_on + timedelta(days=rng.randint(1, 7)) if status in {"delivered", "returned"} else None
                shipments.append({
                    "shipment_id": f"SHP-{seed:04d}-{shipment_counter:07d}",
                    "order_id": order_id,
                    "customer_id": customer["customer_id"],
                    "carrier": rng.choice(CARRIERS),
                    "tracking_number": f"TRK{seed:04d}{shipment_counter:012d}",
                    "status": "delivered" if delivered_on else "in transit",
                    "shipped_on": shipped_on,
                    "delivered_on": delivered_on,
                })

        for _ in range(rng.choices([0, 1, 2, 3], [48, 35, 14, 3])[0]):
            ticket_counter += 1
            opened = earliest + timedelta(days=rng.randrange(date_span + 1))
            tickets.append({
                "ticket_id": f"TKT-{seed:04d}-{ticket_counter:06d}",
                "customer_id": customer["customer_id"],
                "order_id": rng.choice(customer_order_ids) if customer_order_ids and rng.random() < 0.66 else None,
                "opened_at": opened,
                "category": rng.choice(TICKET_CATEGORIES),
                "priority": rng.choices(["low", "normal", "high", "urgent"], [18, 58, 20, 4])[0],
                "status": rng.choices(["resolved", "in progress", "open"], [72, 18, 10])[0],
                "subject": rng.choice(SUBJECTS),
                "document_key": f"{prefix}support/{f'TKT-{seed:04d}-{ticket_counter:06d}'}.docx",
            })

    return {
        "seed": seed,
        "prefix": prefix,
        "customers": customers,
        "products": products,
        "orders": orders,
        "order_items": order_items,
        "invoices": invoices,
        "shipments": shipments,
        "tickets": tickets,
        "tracking_key": f"{prefix}exports/carrier_tracking.csv",
        "catalog_key": f"{prefix}exports/catalog_inventory.xlsx",
    }


def create_local_files(data: dict[str, Any], output: Path) -> dict[str, Path]:
    output.mkdir(parents=True, exist_ok=True)
    prefix = data["prefix"]
    key_to_path: dict[str, Path] = {}
    customer_lookup = {row["customer_id"]: row for row in data["customers"]}
    order_lookup = {row["order_id"]: row for row in data["orders"]}
    product_lookup = {row["sku"]: row["product_name"] for row in data["products"]}

    def destination(key: str) -> Path:
        relative = key.removeprefix(prefix)
        path = output / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        key_to_path[key] = path
        return path

    for customer in data["customers"]:
        path = destination(customer["agreement_key"])
        doc = WordDocument()
        doc.add_heading("Customer Supply Agreement", 0)
        doc.add_paragraph(f"Agreement: AGR-{data['seed']:04d}-{customer['customer_id'][-6:]}")
        doc.add_paragraph(f"Customer: {customer['company_name']} ({customer['customer_id']})")
        doc.add_paragraph(f"Primary contact: {customer['contact_name']} · {customer['email']}")
        doc.add_paragraph(f"Service region: {customer['city']}, {customer['state']} {customer['postal_code']}")
        doc.add_heading("Commercial terms", level=1)
        doc.add_paragraph(
            "Standard net-30 invoicing applies. Orders are subject to product availability; "
            "shipment tracking is provided by the carrier after dispatch. Warranty coverage "
            "follows the product terms supplied with each item."
        )
        doc.add_paragraph("This synthetic agreement is for integration development only.")
        doc.save(path)

    for invoice in data["invoices"]:
        path = destination(invoice["object_key"])
        buffer = io.BytesIO()
        pdf = canvas.Canvas(buffer, pagesize=letter)
        pdf.setTitle(f"Invoice {invoice['invoice_id']}")
        pdf.setFont("Helvetica-Bold", 18)
        pdf.drawString(48, 744, "INVOICE")
        pdf.setFont("Helvetica", 10)
        pdf.drawString(48, 718, f"Invoice: {invoice['invoice_id']}")
        pdf.drawString(48, 702, f"Order: {invoice['order_id']}")
        customer_id = order_lookup[invoice["order_id"]]["customer_id"]
        pdf.drawString(48, 686, f"Customer: {customer_id} · {customer_lookup[customer_id]['company_name'][:40]}")
        pdf.drawString(48, 670, f"Issued: {invoice['issued_on']}    Due: {invoice['due_on']}    Status: {invoice['status']}")
        y = 630
        pdf.setFont("Helvetica-Bold", 10)
        pdf.drawString(48, y, "SKU")
        pdf.drawString(175, y, "Description")
        pdf.drawRightString(430, y, "Qty")
        pdf.drawRightString(505, y, "Unit price")
        pdf.drawRightString(564, y, "Line total")
        y -= 20
        pdf.setFont("Helvetica", 9)
        for item in invoice["items"]:
            if y < 100:
                pdf.showPage()
                y = 744
            pdf.drawString(48, y, item["sku"])
            pdf.drawString(175, y, product_lookup[item["sku"]][:38])
            pdf.drawRightString(430, y, str(item["quantity"]))
            pdf.drawRightString(505, y, f"${item['unit_price']:.2f}")
            pdf.drawRightString(564, y, f"${item['line_total']:.2f}")
            y -= 16
        order = order_lookup[invoice["order_id"]]
        y -= 15
        for label, amount in (("Subtotal", order["subtotal"]), ("Tax", order["tax"]),
                              ("Shipping", order["shipping"]), ("Total", order["total"])):
            pdf.drawRightString(505, y, label)
            pdf.drawRightString(564, y, f"${amount:.2f}")
            y -= 16
        pdf.setFont("Helvetica-Oblique", 8)
        pdf.drawString(48, 48, "Synthetic supplier dataset generated for local integration testing.")
        pdf.save()
        path.write_bytes(buffer.getvalue())

    for ticket in data["tickets"]:
        path = destination(ticket["document_key"])
        customer = customer_lookup[ticket["customer_id"]]
        doc = WordDocument()
        doc.add_heading("Customer Support Case", 0)
        doc.add_paragraph(f"Ticket: {ticket['ticket_id']} · {ticket['status']} · {ticket['priority']} priority")
        doc.add_paragraph(f"Customer: {customer['company_name']} ({ticket['customer_id']})")
        doc.add_paragraph(f"Opened: {ticket['opened_at']} · Category: {ticket['category']}")
        if ticket["order_id"]:
            doc.add_paragraph(f"Related order: {ticket['order_id']}")
        doc.add_heading(ticket["subject"], level=1)
        doc.add_paragraph(
            f"The customer contacted support regarding {ticket['category']}. The service team reviewed "
            "the account and related order records, confirmed the next action with the customer, and "
            "recorded the resolution in the support system. This example case is synthetic."
        )
        doc.save(path)

    tracking_path = destination(data["tracking_key"])
    with tracking_path.open("w", encoding="utf-8", newline="") as file:
        fields = ["shipment_id", "order_id", "customer_id", "carrier", "tracking_number", "status", "shipped_on", "delivered_on"]
        writer = csv.DictWriter(file, fieldnames=fields)
        writer.writeheader()
        for shipment in data["shipments"]:
            writer.writerow({key: iso(value) if isinstance(value, (date, datetime)) else value for key, value in shipment.items()})

    catalog_path = destination(data["catalog_key"])
    workbook = Workbook(write_only=True)
    products_sheet = workbook.create_sheet("Products")
    products_sheet.append(["sku", "product_name", "category", "unit_price", "active", "reorder_level", "on_hand"])
    for product in data["products"]:
        products_sheet.append([product["sku"], product["product_name"], product["category"], float(product["unit_price"]), product["active"], product["reorder_level"], product["on_hand"]])
    orders_sheet = workbook.create_sheet("Order Summary")
    orders_sheet.append(["order_id", "customer_id", "placed_on", "status", "total"])
    for order in data["orders"]:
        orders_sheet.append([order["order_id"], order["customer_id"], iso(order["placed_on"]), order["status"], float(order["total"])])
    workbook.save(catalog_path)

    manifest = {
        "description": "Synthetic B2B equipment supplier data; all records and documents are fictional.",
        "seed": data["seed"],
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "postgres_schema": DB_SCHEMA,
        "s3_prefix": data["prefix"],
        "counts": {key: len(data[key]) for key in ("customers", "products", "orders", "order_items", "invoices", "shipments", "tickets")},
        "object_counts_by_type": {
            "docx_customer_agreements": len(data["customers"]),
            "pdf_invoices": len(data["invoices"]),
            "docx_support_cases": len(data["tickets"]),
            "csv_tracking_exports": 1,
            "xlsx_catalog_and_order_summary": 1,
        },
        "relationships": {
            "customers.agreement_key": "S3 object key",
            "invoices.object_key": "S3 invoice PDF",
            "shipments appear in tracking_export_key": "S3 CSV with shipment, order, and customer IDs",
            "products appear in catalog_key": "S3 Excel workbook; SKUs join to order_items",
            "support_tickets.document_key": "S3 support DOCX",
        },
    }
    manifest_path = output / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    key_to_path[f"{data['prefix']}manifest.json"] = manifest_path
    return key_to_path


def seed_postgres(data: dict[str, Any], config: dict[str, Any]) -> None:
    with psycopg.connect(**config, connect_timeout=10) as connection:
        connection.execute(f"CREATE SCHEMA IF NOT EXISTS {DB_SCHEMA}")
        connection.execute(f"""
            CREATE TABLE IF NOT EXISTS {DB_SCHEMA}.customers (
                customer_id text PRIMARY KEY, company_name text NOT NULL, contact_name text NOT NULL,
                email text NOT NULL UNIQUE, phone text NOT NULL, segment text NOT NULL,
                city text NOT NULL, state text NOT NULL, postal_code text NOT NULL,
                created_at date NOT NULL, agreement_key text NOT NULL
            )
        """)
        connection.execute(f"""
            CREATE TABLE IF NOT EXISTS {DB_SCHEMA}.products (
                sku text PRIMARY KEY, product_name text NOT NULL, category text NOT NULL,
                unit_price numeric(10,2) NOT NULL, active boolean NOT NULL,
                reorder_level integer NOT NULL, on_hand integer NOT NULL,
                catalog_key text NOT NULL
            )
        """)
        connection.execute(f"""
            CREATE TABLE IF NOT EXISTS {DB_SCHEMA}.orders (
                order_id text PRIMARY KEY, customer_id text NOT NULL REFERENCES {DB_SCHEMA}.customers(customer_id),
                placed_on date NOT NULL, status text NOT NULL, subtotal numeric(12,2) NOT NULL,
                tax numeric(12,2) NOT NULL, shipping numeric(12,2) NOT NULL, total numeric(12,2) NOT NULL
            )
        """)
        connection.execute(f"""
            CREATE TABLE IF NOT EXISTS {DB_SCHEMA}.order_items (
                item_id text PRIMARY KEY, order_id text NOT NULL REFERENCES {DB_SCHEMA}.orders(order_id),
                sku text NOT NULL REFERENCES {DB_SCHEMA}.products(sku), quantity integer NOT NULL,
                unit_price numeric(10,2) NOT NULL, line_total numeric(12,2) NOT NULL
            )
        """)
        connection.execute(f"""
            CREATE TABLE IF NOT EXISTS {DB_SCHEMA}.invoices (
                invoice_id text PRIMARY KEY, order_id text NOT NULL UNIQUE REFERENCES {DB_SCHEMA}.orders(order_id),
                issued_on date NOT NULL, due_on date NOT NULL, status text NOT NULL,
                total numeric(12,2) NOT NULL, object_key text NOT NULL
            )
        """)
        connection.execute(f"""
            CREATE TABLE IF NOT EXISTS {DB_SCHEMA}.shipments (
                shipment_id text PRIMARY KEY, order_id text NOT NULL UNIQUE REFERENCES {DB_SCHEMA}.orders(order_id),
                customer_id text NOT NULL REFERENCES {DB_SCHEMA}.customers(customer_id), carrier text NOT NULL,
                tracking_number text NOT NULL UNIQUE, status text NOT NULL, shipped_on date NOT NULL,
                delivered_on date, tracking_export_key text NOT NULL
            )
        """)
        connection.execute(f"""
            CREATE TABLE IF NOT EXISTS {DB_SCHEMA}.support_tickets (
                ticket_id text PRIMARY KEY, customer_id text NOT NULL REFERENCES {DB_SCHEMA}.customers(customer_id),
                order_id text REFERENCES {DB_SCHEMA}.orders(order_id), opened_at date NOT NULL,
                category text NOT NULL, priority text NOT NULL, status text NOT NULL,
                subject text NOT NULL, document_key text NOT NULL
            )
        """)

        # Upsert deterministic IDs so repeated runs refresh this seed without
        # deleting unrelated schemas, rows, or S3 prefixes.
        operations = [
            ("customers", data["customers"], ["customer_id", "company_name", "contact_name", "email", "phone", "segment", "city", "state", "postal_code", "created_at", "agreement_key"]),
            ("products", data["products"], ["sku", "product_name", "category", "unit_price", "active", "reorder_level", "on_hand"]),
            ("orders", data["orders"], ["order_id", "customer_id", "placed_on", "status", "subtotal", "tax", "shipping", "total"]),
            ("order_items", data["order_items"], ["item_id", "order_id", "sku", "quantity", "unit_price", "line_total"]),
            ("invoices", data["invoices"], ["invoice_id", "order_id", "issued_on", "due_on", "status", "total", "object_key"]),
            ("shipments", data["shipments"], ["shipment_id", "order_id", "customer_id", "carrier", "tracking_number", "status", "shipped_on", "delivered_on"]),
            ("support_tickets", data["tickets"], ["ticket_id", "customer_id", "order_id", "opened_at", "category", "priority", "status", "subject", "document_key"]),
        ]
        for table, rows, columns in operations:
            saved_columns = columns[:]
            if table == "products":
                saved_columns.append("catalog_key")
                rows = [dict(row, catalog_key=data["catalog_key"]) for row in rows]
            elif table == "shipments":
                saved_columns.append("tracking_export_key")
                rows = [dict(row, tracking_export_key=data["tracking_key"]) for row in rows]
            if not rows:
                continue
            names = ", ".join(saved_columns)
            placeholders = ", ".join(["%s"] * len(saved_columns))
            updates = ", ".join(f"{column} = EXCLUDED.{column}" for column in saved_columns[1:])
            sql = f"INSERT INTO {DB_SCHEMA}.{table} ({names}) VALUES ({placeholders}) ON CONFLICT ({saved_columns[0]}) DO UPDATE SET {updates}"
            values = [tuple(row[column] for column in saved_columns) for row in rows]
            connection.cursor().executemany(sql, values)


def make_s3_client(config: dict[str, str]):
    return boto3.client(
        "s3",
        aws_access_key_id=config["access_key"],
        aws_secret_access_key=config["secret_key"],
        region_name="us-east-1",
        endpoint_url=config["endpoint"],
    )


def upload_files(client: Any, bucket: str, files: dict[str, Path], workers: int) -> None:
    try:
        client.head_bucket(Bucket=bucket)
    except ClientError as exc:
        status = exc.response.get("ResponseMetadata", {}).get("HTTPStatusCode")
        if status not in (403, 404):
            raise
        client.create_bucket(Bucket=bucket)

    def upload(item: tuple[str, Path]) -> str:
        key, path = item
        client.upload_file(str(path), bucket, key)
        return key

    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = [pool.submit(upload, item) for item in files.items()]
        completed = 0
        for future in as_completed(futures):
            future.result()
            completed += 1
            if completed % 250 == 0 or completed == len(futures):
                print(f"Uploaded {completed:,}/{len(futures):,} objects", flush=True)


def verify_connectors(data: dict[str, Any]) -> None:
    db_config = DatabaseConfig(
        type="postgres",
        host="127.0.0.1",
        port=int(os.getenv("POSTGRES_PORT", "5432")),
        database=need_env("POSTGRES_DB"),
        username=need_env("POSTGRES_USER"),
        password=need_env("POSTGRES_PASSWORD"),
        tables=[f"{DB_SCHEMA}.{name}" for name in ("customers", "orders", "invoices", "shipments", "support_tickets")],
    )
    postgres = PostgresConnector(db_config)
    postgres.check_connection()
    counts = {
        table: postgres.count_rows(f"{DB_SCHEMA}.{table}")
        for table in ("customers", "products", "orders", "order_items", "invoices", "shipments", "support_tickets")
    }
    sample = next(postgres.iter_rows(f"{DB_SCHEMA}.customers", batch_size=1), None)
    print("PostgresConnector read verification:")
    print(json.dumps({"tables": postgres.list_tables(), "counts": counts, "sample_customer": sample}, indent=2, default=str))

    s3_config = S3Config(
        type="s3",
        bucket=BUCKET,
        access_key_id=need_env("MINIO_ROOT_USER"),
        secret_access_key=need_env("MINIO_ROOT_PASSWORD"),
        endpoint_url=f"http://127.0.0.1:{os.getenv('MINIO_API_PORT', '9000')}",
        region="us-east-1",
        prefix=data["prefix"],
    )
    s3 = S3Connector(s3_config)
    s3.check_connection()
    object_count, total_bytes = s3.object_count_and_size()
    sample_objects = [obj.__dict__ for _, obj in zip(range(3), s3.iter_objects())]
    print("S3Connector read verification:")
    print(json.dumps({"bucket": BUCKET, "prefix": data["prefix"], "object_count": object_count, "total_bytes": total_bytes, "first_objects": sample_objects}, indent=2, default=str))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--customers", type=int, default=1000)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--workers", type=int, default=12)
    parser.add_argument("--skip-postgres", action="store_true", help="Generate and upload files without seeding Postgres")
    parser.add_argument("--skip-s3", action="store_true", help="Generate files and seed Postgres without uploading to S3")
    args = parser.parse_args()
    if args.customers < 1 or args.workers < 1:
        parser.error("--customers and --workers must be positive")

    load_env(ENV_FILE)
    print(f"Generating {args.customers:,} customers with random seed {args.seed}...", flush=True)
    data = build_data(args.seed, args.customers)
    output = OUTPUT_ROOT / f"seed-{args.seed:04d}"
    files = create_local_files(data, output)
    print(f"Generated {len(files):,} files in {output}", flush=True)

    if not args.skip_postgres:
        postgres_config = {
            "host": "127.0.0.1",
            "port": int(os.getenv("POSTGRES_PORT", "5432")),
            "dbname": need_env("POSTGRES_DB"),
            "user": need_env("POSTGRES_USER"),
            "password": need_env("POSTGRES_PASSWORD"),
        }
        seed_postgres(data, postgres_config)
        print(f"Seeded Postgres schema {DB_SCHEMA}.", flush=True)

    if not args.skip_s3:
        s3_config = {
            "access_key": need_env("MINIO_ROOT_USER"),
            "secret_key": need_env("MINIO_ROOT_PASSWORD"),
            "endpoint": f"http://127.0.0.1:{os.getenv('MINIO_API_PORT', '9000')}",
        }
        upload_files(make_s3_client(s3_config), BUCKET, files, args.workers)
        print(f"Uploaded under s3://{BUCKET}/{data['prefix']}", flush=True)

    if not args.skip_postgres and not args.skip_s3:
        verify_connectors(data)

    counts = {name: len(data[name]) for name in ("customers", "products", "orders", "order_items", "invoices", "shipments", "tickets")}
    print("\nDataset summary:")
    print(json.dumps({"seed": args.seed, "counts": counts, "generated_files": len(files), "local_path": str(output)}, indent=2))


if __name__ == "__main__":
    main()
