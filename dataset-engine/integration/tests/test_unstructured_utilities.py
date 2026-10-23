from __future__ import annotations

import io
from contextlib import contextmanager

import pytest
from docx import Document as WordDocument
from reportlab.pdfgen import canvas

from main.connect.object_storage.s3 import SourceObject
from main.document.unstructured_data.documents import (
    iter_unstructured_documents,
    normalize_text,
    text_to_document,
)
from main.schema.document import SourceType


class TrackedBytesIO(io.BytesIO):
    def __init__(self, content: bytes) -> None:
        super().__init__(content)
        self.closed_by_connector = False

    def close(self) -> None:
        self.closed_by_connector = True
        super().close()


class MemoryObjectConnector:
    """S3Connector-shaped test double that tracks streaming body cleanup."""

    def __init__(self, objects: dict[str, bytes]) -> None:
        self.objects = objects
        self.bodies: list[TrackedBytesIO] = []

    @contextmanager
    def open_object(self, key: str):
        body = TrackedBytesIO(self.objects[key])
        self.bodies.append(body)
        try:
            yield body
        finally:
            body.close()


def source(key: str) -> SourceObject:
    return SourceObject(key=key, size=0)


def test_normalize_text_strips_bom_crlf_trailing_space_and_extra_blank_lines():
    assert normalize_text("\ufeff# Title\r\n\r\n\r\nBody line  \r\n") == (
        "# Title\n\nBody line"
    )


def test_text_document_identity_is_stable_while_signature_tracks_content():
    original = text_to_document(
        "# Notes\n\nfirst version",
        namespace="supplier",
        source_id="notes.md",
        metadata={"extension": ".md"},
    )
    updated = text_to_document(
        "# Notes\n\nsecond version",
        namespace="supplier",
        source_id="notes.md",
        metadata={"extension": ".md"},
    )
    elsewhere = text_to_document(
        "# Notes\n\nfirst version",
        namespace="other",
        source_id="notes.md",
    )

    assert original.source_type == SourceType.S3_OBJECT
    assert original.doc_id == updated.doc_id
    assert original.doc_id != elsewhere.doc_id
    assert original.signature != updated.signature
    assert original.raw_content == "# Notes\n\nfirst version"
    assert original.metadata["namespace"] == "supplier"
    assert original.metadata["source_id"] == "notes.md"
    assert original.metadata["word_count"] == 4


def test_markdown_reader_normalizes_text_and_closes_stream():
    connector = MemoryObjectConnector({
        "notes.md": b'\xef\xbb\xbf# Title\r\n\r\n\r\nBody line  \r\n',
    })

    documents = list(iter_unstructured_documents(
        connector, source("notes.md"), namespace="supplier",
    ))

    assert len(documents) == 1
    assert documents[0].raw_content == "# Title\n\nBody line"
    assert documents[0].metadata["extension"] == ".md"
    assert documents[0].metadata["media_type"] == "text/markdown"
    assert connector.bodies[0].closed_by_connector


@pytest.mark.parametrize("content", [b"", b"  \r\n\r\n\t"])
def test_empty_documents_are_skipped(content: bytes):
    connector = MemoryObjectConnector({"empty.txt": content})
    assert list(iter_unstructured_documents(connector, source("empty.txt"))) == []
    assert connector.bodies[0].closed_by_connector


def test_docx_reader_preserves_headings_lists_and_table_order():
    document = WordDocument()
    document.core_properties.title = "Customer Agreement"
    document.core_properties.author = "Sab"
    document.add_heading("Customer Agreement", level=0)
    document.add_paragraph("Agreement: AGR-1")
    document.add_heading("Commercial terms", level=1)
    document.add_paragraph("Net-30 invoicing", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Days"
    table.cell(1, 1).text = "30"
    document.add_paragraph("Final note")

    buffer = io.BytesIO()
    document.save(buffer)

    connector = MemoryObjectConnector({"agreement.docx": buffer.getvalue()})
    documents = list(iter_unstructured_documents(
        connector, source("agreement.docx"), namespace="supplier",
    ))

    assert len(documents) == 1
    text = documents[0].raw_content
    assert text.index("# Customer Agreement") < text.index("Agreement: AGR-1")
    assert "# Commercial terms" in text
    assert "- Net-30 invoicing" in text
    assert "| Term | Value |" in text
    assert "| --- | --- |" in text
    assert "| Days | 30 |" in text
    assert text.index("| Days | 30 |") < text.index("Final note")
    assert documents[0].metadata["extension"] == ".docx"
    assert documents[0].metadata["title"] == "Customer Agreement"
    assert documents[0].metadata["author"] == "Sab"
    assert documents[0].metadata["paragraph_count"] == 5
    assert documents[0].metadata["table_count"] == 1
    assert connector.bodies[0].closed_by_connector


def test_pdf_reader_extracts_text_from_every_page():
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(72, 720, "Page one content")
    pdf.showPage()
    pdf.drawString(72, 720, "Page two content")
    pdf.save()

    connector = MemoryObjectConnector({"invoice.pdf": buffer.getvalue()})
    documents = list(iter_unstructured_documents(
        connector, source("invoice.pdf"), namespace="supplier",
    ))

    assert len(documents) == 1
    assert "Page one content" in documents[0].raw_content
    assert "Page two content" in documents[0].raw_content
    assert documents[0].metadata["extension"] == ".pdf"
    assert documents[0].metadata["page_count"] == 2
    assert connector.bodies[0].closed_by_connector


def test_dispatch_rejects_unsupported_extensions():
    connector = MemoryObjectConnector({})
    with pytest.raises(ValueError, match="Unsupported unstructured file"):
        list(iter_unstructured_documents(connector, source("image.png")))
